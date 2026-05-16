"""Storage layer: Excel (openpyxl), Notes.txt, and JSON settings persistence."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

DATA_DIR = Path.home() / "jobsoffers"
EXCEL_PATH = DATA_DIR / "JobsData.xlsx"
NOTES_PATH = DATA_DIR / "Notes.txt"
SETTINGS_PATH = DATA_DIR / "settings.json"
HEADER_FILL = PatternFill(start_color="0F3460", end_color="0F3460", fill_type="solid")
HEADER_FONT = Font(color="EAEAEA", bold=True, size=11)
ALT_FILL = PatternFill(start_color="16213E", end_color="16213E", fill_type="solid")

JOB_FIELDS = [
    "Status", "Title", "Company", "Location", "Source",
    "URL", "Match_Score", "Languages", "Date_Found", "Notes",
]


def ensure_data_dir() -> None:
    """Create ~/jobsoffers/ and default files if missing."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not EXCEL_PATH.exists():
        _create_empty_excel()
    if not NOTES_PATH.exists():
        NOTES_PATH.write_text("", encoding="utf-8")
    if not SETTINGS_PATH.exists():
        _save_settings(_default_settings())


def _create_empty_excel() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Jobs"
    for col_idx, field in enumerate(JOB_FIELDS, 1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 50
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 20
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 30
    ws.freeze_panes = "A2"
    wb.save(str(EXCEL_PATH))


def _default_settings() -> dict[str, Any]:
    return {
        "gmail_email": "",
        "gmail_app_password": "",
        "keywords": [],
        "locations": {"mexico": True, "global_remote": True, "arabic_remote": True},
        "languages": {"arabic": False, "english": True, "french": False, "spanish": False},
        "min_match_score": 30,
        "schedule_time": "08:00",
    }


def load_settings() -> dict[str, Any]:
    ensure_data_dir()
    try:
        data = json.loads(SETTINGS_PATH.read_text(encoding="utf-8"))
        defaults = _default_settings()
        defaults.update(data)
        return defaults
    except (json.JSONDecodeError, FileNotFoundError):
        return _default_settings()


def _save_settings(settings: dict[str, Any]) -> None:
    SETTINGS_PATH.write_text(json.dumps(settings, indent=2), encoding="utf-8")


def save_settings(settings: dict[str, Any]) -> None:
    ensure_data_dir()
    _save_settings(settings)


def add_jobs(jobs: list[dict[str, Any]]) -> int:
    """Append jobs to Excel. Returns count of new rows."""
    ensure_data_dir()
    wb = openpyxl.load_workbook(str(EXCEL_PATH))
    ws = wb.active
    existing_urls: set[str] = set()
    for row in range(2, ws.max_row + 1):
        url = ws.cell(row=row, column=6).value or ""
        existing_urls.add(url.strip())

    added = 0
    for job in jobs:
        url = (job.get("url") or "").strip()
        if url and url in existing_urls:
            continue
        row = ws.max_row + 1
        ws.cell(row=row, column=1, value=job.get("status", "New"))
        ws.cell(row=row, column=2, value=job.get("title", ""))
        ws.cell(row=row, column=3, value=job.get("company", ""))
        ws.cell(row=row, column=4, value=job.get("location", ""))
        ws.cell(row=row, column=5, value=job.get("source", ""))
        ws.cell(row=row, column=6, value=url)
        ws.cell(row=row, column=7, value=job.get("match_score", 0))
        ws.cell(row=row, column=8, value=", ".join(job.get("languages", [])))
        ws.cell(row=row, column=9, value=datetime.now().strftime("%Y-%m-%d"))
        ws.cell(row=row, column=10, value=job.get("notes", ""))
        if row % 2 == 0:
            for col in range(1, len(JOB_FIELDS) + 1):
                ws.cell(row=row, column=col).fill = ALT_FILL
        if url:
            existing_urls.add(url)
        added += 1

    wb.save(str(EXCEL_PATH))
    return added


def get_all_jobs() -> list[dict[str, Any]]:
    ensure_data_dir()
    wb = openpyxl.load_workbook(str(EXCEL_PATH))
    ws = wb.active
    jobs: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        jobs.append({
            "status": ws.cell(row=row, column=1).value or "",
            "title": ws.cell(row=row, column=2).value or "",
            "company": ws.cell(row=row, column=3).value or "",
            "location": ws.cell(row=row, column=4).value or "",
            "source": ws.cell(row=row, column=5).value or "",
            "url": ws.cell(row=row, column=6).value or "",
            "match_score": ws.cell(row=row, column=7).value or 0,
            "languages": ws.cell(row=row, column=8).value or "",
            "date_found": ws.cell(row=row, column=9).value or "",
            "notes": ws.cell(row=row, column=10).value or "",
        })
    return jobs


def update_job_status(url: str, new_status: str) -> None:
    ensure_data_dir()
    wb = openpyxl.load_workbook(str(EXCEL_PATH))
    ws = wb.active
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=6).value == url:
            ws.cell(row=row, column=1, value=new_status)
            break
    wb.save(str(EXCEL_PATH))


def append_note(text: str) -> None:
    ensure_data_dir()
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    with open(str(NOTES_PATH), "a", encoding="utf-8") as f:
        f.write(f"[{timestamp}] {text}\n")


def get_notes() -> str:
    ensure_data_dir()
    return NOTES_PATH.read_text(encoding="utf-8")