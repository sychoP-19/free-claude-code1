"""SKILL: docx — Generate speaker notes and scripts as Word docs."""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def write_speaker_notes(text: str, output_path: str, title: str = "Speaker Notes"):
    doc = Document()
    doc.title = title

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for paragraph_text in text.split("\n\n"):
        p = doc.add_paragraph(paragraph_text.strip())
        p.paragraph_format.space_after = Pt(6)

    doc.save(output_path)
