"""Layer 3 — Supporting Skills: file-reading, xlsx, pdf-reading.

Invoked on-demand when users upload source material
(brief PDF, data CSV, existing deck, etc.).
"""

from __future__ import annotations

from pathlib import Path


class FileReaderSkill:
    def __init__(self, config: dict):
        self.config = config

    def read(self, file_paths: list[str]) -> dict:
        result = {"ok": True, "extra_context": "", "parsed": {}}

        for fp in file_paths:
            p = Path(fp)
            if not p.exists():
                continue

            suffix = p.suffix.lower()
            if suffix == ".pdf":
                content = self._read_pdf(p)
            elif suffix in (".xlsx", ".xls", ".csv"):
                content = self._read_spreadsheet(p)
            elif suffix in (".pptx", ".ppt"):
                content = self._read_pptx(p)
            elif suffix in (".docx", ".doc"):
                content = self._read_docx(p)
            elif suffix in (".txt", ".md"):
                content = p.read_text(encoding="utf-8", errors="replace")
            else:
                content = f"[Unsupported format: {suffix}]"

            result["parsed"][p.name] = content[:8000]  # Truncate per file
            result["extra_context"] += f"\n--- {p.name} ---\n{content[:4000]}\n"

        return result

    def _read_pdf(self, path: Path) -> str:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages[:50])
        except ImportError:
            return "[PyPDF2 not installed]"

    def _read_spreadsheet(self, path: Path) -> str:
        try:
            if path.suffix.lower() == ".csv":
                return path.read_text(encoding="utf-8", errors="replace")[:8000]
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True)
            rows = []
            for ws in wb.worksheets[:5]:
                for row in ws.iter_rows(max_row=100, values_only=True):
                    rows.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(rows)[:8000]
        except ImportError:
            return "[openpyxl not installed]"

    def _read_pptx(self, path: Path) -> str:
        try:
            from pptx import Presentation
            prs = Presentation(str(path))
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text_frame.text)
            return "\n".join(texts)[:8000]
        except ImportError:
            return "[python-pptx not installed]"

    def _read_docx(self, path: Path) -> str:
        try:
            from docx import Document
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)[:8000]
        except ImportError:
            return "[python-docx not installed]"
