"""SKILL: pptx — Build slide decks with branded overlays."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pathlib import Path


def build_slide_deck(segments: list[dict], output_path: str, title: str = "Presentation", brand_logo: str | None = None):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # blank

    # Title slide
    slide = prs.slides.add_slide(blank_layout)
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Content slides
    for seg in segments:
        slide = prs.slides.add_slide(blank_layout)
        # Narration text (positioned for avatar layout: bottom portion)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(5), Inches(12), Inches(2))
        tf = box.text_frame
        tf.word_wrap = True
        seg_text = seg.get("text", "")
        if len(seg_text) > 200:
            seg_text = seg_text[:200] + "..."
        p = tf.paragraphs[0]
        p.text = seg_text
        p.font.size = Pt(18)
        p.alignment = PP_ALIGN.LEFT

        # Segment label
        label_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(3), Inches(0.5))
        lp = label_box.text_frame.paragraphs[0]
        lp.text = seg.get("id", "")
        lp.font.size = Pt(12)
        lp.font.italic = True

    # Logo overlay
    if brand_logo and Path(brand_logo).exists():
        for slide in prs.slides:
            slide.shapes.add_picture(brand_logo, Inches(11.5), Inches(0.2), Inches(1.5))

    prs.save(output_path)
